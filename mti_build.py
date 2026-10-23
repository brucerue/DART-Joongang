"""무역협회 K-stat의 'HSK-MTI 연계표'(엑셀)를 앱이 읽는 mti_map.json으로 바꾼다 (연계표가 새로 나올 때마다 다시 실행).

- 연계표 시트 '연도별 HSK-MTI 연계표': 기준년도별(2021년 이전 / 2022~2024 / 2025~2026) HS 10자리(HSK) -> MTI 6자리 연결.
- 시트 'MTI 코드 변경 내역': 1·2·3·4·6단위 MTI 이름(2022년 이후 코드, 삭제된 코드는 뺌).
결과: {"source", "tables": [{"label", "from", "to", "mti": {MTI6: [HSK, ...]}, "names": {MTI6: 이름}}], "names": {MTI: 이름}}
사용: python mti_build.py 연계표.xlsx [--out mti_map.json]
"""
import argparse
import json
import re
import sys

import openpyxl


def year_range(label):
    """'2021년 이전' -> (0, 2021), '2022년-2024년' -> (2022, 2024)"""
    years = [int(y) for y in re.findall(r"(\d{4})", label)]
    if "이전" in label:
        return 0, years[0]
    return years[0], years[-1]


def build(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    tables = {}
    for r in wb["연도별 HSK-MTI 연계표"].iter_rows(min_row=2, values_only=True):
        label, hsk, _hs_name, mti, mti_name = (list(r) + [None] * 5)[:5]
        if not (label and hsk and mti):
            continue
        t = tables.setdefault(str(label).strip(), {"mti": {}, "names": {}})
        t["mti"].setdefault(str(mti).strip(), []).append(str(hsk).strip())
        t["names"][str(mti).strip()] = str(mti_name or "").strip()
    names = {}
    for r in wb["MTI 코드 변경 내역"].iter_rows(min_row=2, values_only=True):
        if r[0] is None or r[4] == "삭제":
            continue
        names[str(r[0]).strip()] = str(r[2] or "").strip()
    out_tables = []
    for label, t in sorted(tables.items(), key=lambda kv: year_range(kv[0])):
        lo, hi = year_range(label)
        out_tables.append({"label": label, "from": lo, "to": hi, "mti": {k: sorted(v) for k, v in sorted(t["mti"].items())}, "names": t["names"]})
    return {"source": "무역협회 K-stat HSK-MTI 연계표", "tables": out_tables, "names": names}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("--out", default="mti_map.json")
    args = ap.parse_args()
    data = build(args.xlsx)
    with open(args.out, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, separators=(",", ":"))
    print(f"표 {[(t['label'], len(t['mti'])) for t in data['tables']]}, 이름 {len(data['names'])}개")


if __name__ == "__main__":
    sys.exit(main())
